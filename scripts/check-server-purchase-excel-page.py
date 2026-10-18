"""Complete only the missing purchase Excel and print-template page checks.

Normal API/browser writes; read-only SQL. Existing purchase data stays unchanged.
No output copies, screenshots, build, deployment or completed-case replay.
"""
import base64
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import time

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("ordinary", ROOT / "scripts/check-server-remaining-functional.py")
ordinary = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ordinary)


def workbook_source():
    book = Workbook()
    sheet = book.active
    sheet.title = "采购单"
    sheet.append(["采购 Excel 验收"])
    sheet.merge_cells("A1:F1")
    sheet.append(["采购单号", "{{purchaseNo}}", None, None, None, None, None, 1.1])
    sheet.append(["物品名称", None, "数量", "单价", "金额", "固定系数后金额"])
    sheet.merge_cells("A3:B3")
    sheet.append(["{{#items}}{{consumableName}}", None, "{{quantity}}", "{{unitPrice}}", "=C4*D4", "=E4*$H$2"])
    sheet.merge_cells("A4:B4")
    sheet.append(["合计", None, None, None, "=SUM(E4:E4)"])
    sheet.row_dimensions[4].height = 32
    edge = Side(style="thin", color="334155")
    for row in sheet.iter_rows(min_row=3, max_row=4, min_col=1, max_col=6):
        for cell in row:
            cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row == 3:
                cell.font = Font(name="宋体", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="164E63")
            else:
                cell.font = Font(name="宋体", color="164E63")
                cell.fill = PatternFill("solid", fgColor="E0F2FE")
                if cell.column >= 4:
                    cell.number_format = "0.00"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 10
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.print_area = "A1:F5"
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def verify_workbook(payload, purchase_no, items):
    book = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        sheet = book["采购单"]
        assert sheet["B2"].value == purchase_no, "Actual purchase number must be filled"
        assert str(sheet["E" + str(4 + len(items))].value) == f"=SUM(E4:E{3 + len(items)})", "Expanded summary formula"
        merges = {str(value) for value in sheet.merged_cells.ranges}
        assert {"A1:F1", "A3:B3"}.issubset(merges), "Original title/header merges"
        assert sheet["A3"].font.bold and sheet["A3"].fill.fgColor.rgb.endswith("164E63"), "Original header style"
        for index, item in enumerate(items, 4):
            assert sheet[f"A{index}"].value == item["consumableName"]
            assert sheet[f"C{index}"].value == item["quantity"]
            assert float(sheet[f"D{index}"].value) == float(item["unitPrice"])
            assert sheet[f"E{index}"].value == f"=C{index}*D{index}"
            assert sheet[f"F{index}"].value == f"=E{index}*$H$2"
            assert f"A{index}:B{index}" in merges
            assert sheet[f"D{index}"].number_format == "0.00"
            assert sheet[f"A{index}"].border.left.style == "thin"
            assert sheet[f"A{index}"].fill.fgColor.rgb.endswith("E0F2FE")
            assert sheet.row_dimensions[index].height == 32
        assert not any("{{" in str(cell.value) for row in sheet for cell in row), "No unresolved variables"
        return {"sha256": hashlib.sha256(payload).hexdigest(), "bytes": len(payload), "items": len(items),
            "sourceValues": True, "typedNumbers": True, "style": True, "formulas": True, "merges": True}
    finally:
        book.close()


def main():
    state_path = ROOT / ".cloudflow-runtime/server-deployment-20261018/purchase-excel-page-20261018.json"
    state = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
    if state.get("complete"):
        print("SKIP completed purchase Excel/template page checks", flush=True)
        return 0

    def save():
        state_path.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")

    checks = ordinary.Checks()
    try:
        checks.account(verify_reset=False, reuse=True)
        assert state.get("operatorId") in (None, checks.user_id)
        code = checks.tag + "_XLSX_PAGE"
        state.update(operatorId=checks.user_id, templateCode=code, templateName="采购 Excel 页面验收")
        save()
        purchase = ordinary.api(checks.owned, "GET", "/oa/purchase/request/1749051")
        assert purchase["status"] == "DRAFT" and purchase["items"]
        page_case = json.loads((state_path.parent / "print-pages-20261018.json").read_text(encoding="utf-8"))["cases"]["PURCHASE_ORDER"]
        assert page_case["bizId"] == purchase["id"] and page_case["complete"]
        rows = checks.sql("SELECT template_id,version,remark FROM sys_print_template WHERE tenant_id=100000 "
            "AND template_code=%s AND deleted=0", (code,))
        assert len(rows) <= 1
        if rows:
            assert state.get("templateId") in (None, rows[0]["template_id"])
            state["templateId"] = rows[0]["template_id"]
            uploaded = ordinary.request(checks.owned, "GET", f"/auth/system/print/template/{state['templateId']}/file")
            assert uploaded.status_code == 200 and uploaded.content.startswith(b"PK")
            source = uploaded.content
            if rows[0]["remark"] == code + " 页面编辑已核对":
                state["edited"] = True
            save()
        else:
            assert "templateId" not in state, "Recorded template is missing; inspect instead of recreating"
            source = workbook_source()
        sample = {"purchaseNo": "PREVIEW-2-ROWS", "items": [
            {"consumableName": "示例纸张", "quantity": 2, "unitPrice": 12.5},
            {"consumableName": "示例墨盒", "quantity": 3, "unitPrice": 24}]}

        def browser(script, settings, records_only=False):
            settings.update(entry=ordinary.support.ENTRY, cookies=[
                {"name": cookie.name, "value": cookie.value, "domain": ordinary.support.HOST,
                 "path": cookie.path or "/", "httpOnly": cookie.name == "Authorization",
                 "secure": False, "sameSite": "Lax"} for cookie in checks.owned.cookies])
            result = subprocess.run(["node", "scripts/" + script], cwd=ROOT / "cloudflow-frontend",
                input=json.dumps(settings), text=True, encoding="utf-8", capture_output=True, timeout=90)
            for line in result.stdout.splitlines():
                if " " not in line:
                    print(line, flush=True)
                    continue
                event, serialized = line.split(" ", 1)
                if event in {"TEMPLATE_CREATED", "TEMPLATE_VIEWED", "TEMPLATE_EDITED", "TEMPLATE_PREVIEWED",
                             "TEMPLATE_RECORDS", "TEMPLATE_PAGE_PASS", "BROWSER_PRINT"}:
                    value = json.loads(serialized)
                    workbook = base64.b64decode(value.pop("workbook")) if "workbook" in value else None
                    if event == "TEMPLATE_CREATED":
                        state["templateId"] = value["templateId"]
                    elif event == "TEMPLATE_VIEWED":
                        state["sourceViewed"] = value["uploadedSourceIdentical"]
                    elif event == "TEMPLATE_EDITED":
                        state.update(edited=True, version=value["version"])
                    elif event == "TEMPLATE_PREVIEWED":
                        value["validated"] = verify_workbook(workbook, sample["purchaseNo"], sample["items"])
                        state.update(previewed=True, preview=value)
                    elif event == "TEMPLATE_PAGE_PASS" and not records_only:
                        state["managementBrowserComplete"] = True
                    elif event == "TEMPLATE_RECORDS":
                        state["records"] = value
                    elif event == "BROWSER_PRINT":
                        value["validated"] = verify_workbook(workbook, purchase["purchaseNo"], purchase["items"])
                        state["browser"] = value
                    save()
                    print(event, json.dumps(value, ensure_ascii=False), flush=True)
                else:
                    print(line, flush=True)
            if result.returncode:
                print(result.stderr[-3000:], flush=True)
                raise AssertionError("Selected browser check failed; completed stages have been saved")

        if not state.get("templateVerified"):
            if not state.get("managementBrowserComplete"):
                state.setdefault("previewRecordCountBefore", checks.sql("SELECT COUNT(*) n FROM sys_print_record "
                    "WHERE tenant_id=100000 AND template_code=%s", (code,))[0]["n"])
                save()
                browser("check-server-print-template-browser.mjs", {**state,
                    "upload": base64.b64encode(source).decode(), "sample": sample})
            checks.check(state.get("previewed") and state.get("edited"), "模板页面上传、查询、原件下载、编辑和示例下载通过")
            record_after = checks.sql("SELECT COUNT(*) n FROM sys_print_record WHERE tenant_id=100000 "
                "AND template_code=%s", (code,))[0]["n"]
            checks.check(state["previewRecordCountBefore"] == record_after, "示例预览不新增业务打印记录")
            template = ordinary.api(checks.owned, "GET", f"/auth/system/print/template/{state['templateId']}")
            checks.check(template["version"] == 1 and template["remark"] == code + " 页面编辑已核对"
                and template["templateFileUrl"], "页面编辑版本递增并保留原Excel附件")
            state["templateVerified"] = True
            save()
        if "browser" not in state:
            records = checks.sql("SELECT record_id FROM sys_print_record WHERE tenant_id=100000 AND template_id=%s",
                (state["templateId"],))
            assert not records, "Print exists without a completed browser result; inspect before replay"
            browser("check-server-functional-browser.mjs", {"printOnly": True, "printCase": {
                "bizType": "PURCHASE_ORDER", "bizId": purchase["id"], "route": "/office/purchase-request",
                "documentNo": purchase["purchaseNo"], "pageNum": page_case["pageNum"],
                "templateName": state["templateName"], "outputFormat": "EXCEL"}})
        records = checks.sql("SELECT record_id,file_size,operator_id,print_time,ip_address FROM sys_print_record "
            "WHERE tenant_id=100000 AND template_id=%s AND biz_type='PURCHASE_ORDER' AND biz_id=%s "
            "AND output_format='EXCEL'", (state["templateId"], str(purchase["id"])))
        checks.check(len(records) == 1 and records[0]["operator_id"] == checks.user_id
            and records[0]["file_size"] == state["browser"]["bytes"] and records[0]["print_time"]
            and records[0]["ip_address"], "采购Excel打印唯一记录与操作者、字节、时间和IP一致")
        state["recordId"] = records[0]["record_id"]
        save()
        if "records" not in state:
            browser("check-server-print-template-browser.mjs", {**state, "recordsOnly": True, "record": {
                "bizId": purchase["id"], "operatorId": checks.user_id, "recordId": state["recordId"],
                "bytes": state["browser"]["bytes"]}}, records_only=True)
        checks.check(purchase == ordinary.api(checks.owned, "GET", "/oa/purchase/request/1749051"),
            "采购原草稿及明细保持一致，没有提交、审批或付款")
        state.update(complete=True, verifiedAt=time.strftime("%Y-%m-%d %H:%M:%S %z"))
        save()
        print("PURCHASE_EXCEL_PAGE_PASS", json.dumps({key: state[key] for key in
            ("operatorId", "templateId", "recordId", "verifiedAt")}), flush=True)
        return 0
    finally:
        checks.close()


if __name__ == "__main__":
    raise SystemExit(main())
