"""Small real imports for businesses still missing ordinary evidence.

All writes use business HTTP APIs; SQL is read only. Successful cases are saved
and skipped on resume. No cleanup, performance run or external integration.
Credentials and tokens stay in process memory. Input workbooks stay in memory.
"""
import argparse
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import secrets
import time

import openpyxl
import pymysql

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("ordinary", ROOT / "scripts/check-server-remaining-functional.py")
ordinary = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ordinary)
api = ordinary.api
request = ordinary.request

BUSINESSES = ("ASSET", "SUPPLIER", "EMPLOYEE", "OPPORTUNITY", "HOLIDAY")


def workbook(headers, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "导入数据"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class Checks:
    def __init__(self, state_path):
        self.path = state_path
        if state_path.exists():
            self.state = json.loads(state_path.read_text(encoding="utf-8"))
        else:
            self.state = {"tag": "IM" + time.strftime("%m%d%H%M%S") + secrets.token_hex(2),
                          "cases": {}, "passed": []}
        self.connection = pymysql.connect(host=ordinary.support.HOST, user="root",
            password=os.environ["CLOUDFLOW_DB_PASSWORD"], database="cloud_flow_db", charset="utf8mb4",
            autocommit=True, connect_timeout=10, read_timeout=15, cursorclass=pymysql.cursors.DictCursor)
        self.client = ordinary.session()
        self.upload_times = {}

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.state, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    def check(self, value, label):
        if not value:
            raise AssertionError(label)
        if label not in self.state["passed"]:
            self.state["passed"].append(label)
        self.save()
        print("PASS", label, flush=True)

    def sql(self, sql, params=()):
        return ordinary.support.query(self.connection, sql, params)

    def read(self, service, task_id):
        return api(self.client, "GET", f"/{service}/system/import/{task_id}")

    def wait(self, service, task_id):
        deadline = time.monotonic() + 35
        while True:
            task = self.read(service, task_id)
            if task["status"] not in ("PENDING", "VALIDATING", "IMPORTING"):
                return task
            if time.monotonic() >= deadline:
                raise AssertionError(f"{service}/{task_id}: import still running after 35 seconds")
            time.sleep(.5)

    def prepare(self, service, biz, key, headers, rows):
        case = self.state["cases"].setdefault(key, {})
        if "taskId" not in case:
            endpoint = (service, biz)
            delay = 5.2 - (time.monotonic() - self.upload_times.get(endpoint, float("-inf")))
            if delay > 0:
                time.sleep(delay)
            task = api(self.client, "POST", f"/{service}/system/import/{biz}/upload",
                files={"file": (self.state["tag"] + "_" + key + ".xlsx", workbook(headers, rows),
                                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            self.upload_times[endpoint] = time.monotonic()
            case["taskId"] = task["taskId"]
            self.save()
            self.check(task["status"] == "UPLOADED" and task["totalRows"] == len(rows), key + "_uploaded")
        task_id = case["taskId"]
        task = self.wait(service, task_id)
        if task["status"] == "UPLOADED" or (task["status"] == "FAILED"
                and task["successRows"] == 0 and task["failRows"] == 0):
            api(self.client, "POST", f"/{service}/system/import/{task_id}/validate")
            task = self.wait(service, task_id)
        case["lastTask"] = task
        self.save()
        return case, task

    def execute(self, service, case, skip=False):
        task_id = case["taskId"]
        task = self.wait(service, task_id)
        if task["status"] == "VALIDATED":
            api(self.client, "POST", f"/{service}/system/import/{task_id}/execute", json={"skipInvalidRows": skip})
            task = self.wait(service, task_id)
        case["lastTask"] = task
        self.save()
        return task

    def result(self, service, biz, case, expected="SUCCESS", total=1, success=1, fail=0):
        task = self.wait(service, case["taskId"])
        self.check(task["status"] == expected and task["totalRows"] == total
                   and task["successRows"] == success and task["failRows"] == fail, biz + "_result_counts")
        self.check(task["bizType"] == biz and task["operatorId"] == self.state["operatorId"]
                   and task["operatorName"] and task["jobLogId"] and task["startTime"] and task["endTime"],
                   biz + "_operator_and_async_job")
        deadline = time.monotonic() + 15
        while True:
            audits = self.sql("SELECT audit_id,create_by,after_json FROM sys_audit_log WHERE tenant_id=100000 "
                "AND audit_field='importResult' AND JSON_EXTRACT(after_json,'$.taskId')=%s", (case["taskId"],))
            notices = self.sql("SELECT record_id,receiver_id,send_status FROM sys_message_record WHERE tenant_id=100000 "
                "AND biz_type='IMPORT_TASK' AND biz_id=%s AND receiver_id=%s", (str(case["taskId"]), self.state["operatorId"]))
            if audits and any(row["send_status"] == "SUCCESS" for row in notices):
                break
            if time.monotonic() >= deadline:
                break
            time.sleep(.5)
        summaries = [json.loads(row["after_json"]) for row in audits]
        self.check(any(value.get("bizType") == biz and value.get("status") == expected
                       and value.get("successRows") == success and value.get("failRows") == fail
                       and value.get("fileName") == task["fileName"] for value in summaries)
                   and all(row["create_by"] == task["operatorName"] for row in audits), biz + "_audit")
        self.check(any(row["send_status"] == "SUCCESS" for row in notices), biz + "_completion_notice")
        case.update(auditIds=[row["audit_id"] for row in audits], noticeIds=[row["record_id"] for row in notices])
        self.save()

    def default_calendar_snapshot(self):
        calendars = self.sql("SELECT * FROM sys_holiday_calendar WHERE tenant_id=100000 "
            "AND default_flag=1 AND deleted=0 ORDER BY calendar_id")
        dates = self.sql("SELECT d.* FROM sys_holiday_date d JOIN sys_holiday_calendar c "
            "ON c.calendar_id=d.calendar_id AND c.tenant_id=d.tenant_id WHERE c.tenant_id=100000 "
            "AND c.default_flag=1 AND c.deleted=0 ORDER BY d.calendar_id,d.holiday_date")
        digest = hashlib.sha256(json.dumps([calendars, dates], ensure_ascii=False,
            sort_keys=True, default=str).encode("utf-8")).hexdigest()
        return {"calendarIds": [row["calendar_id"] for row in calendars],
                "dateCount": len(dates), "sha256": digest}

    def holidays(self):
        case = self.state["cases"].setdefault("HOLIDAY", {})
        if "defaultBaseline" not in case:
            case["defaultBaseline"] = self.default_calendar_snapshot()
            self.check(bool(case["defaultBaseline"]["calendarIds"])
                       and case["defaultBaseline"]["dateCount"] == 28, "HOLIDAY_official_baseline_present")
        code = self.state["tag"] + "_CAL"
        calendars = self.sql("SELECT calendar_id,default_flag,status FROM sys_holiday_calendar "
            "WHERE tenant_id=100000 AND calendar_code=%s AND deleted=0", (code,))
        if not calendars:
            api(self.client, "POST", "/auth/system/holiday/calendar", json={
                "calendarCode": code, "calendarName": "导入验收专用2026日历" + self.state["tag"],
                "region": "CN", "defaultFlag": 0, "weeklyRestDays": "6,7", "status": "ENABLED",
                "workStartTime": "09:00:00", "workEndTime": "18:00:00",
                "lunchStartTime": "12:00:00", "lunchEndTime": "13:00:00",
                "remark": "独立导入验收日历，不关联部门或员工，不替代正式默认日历"})
            calendars = self.sql("SELECT calendar_id,default_flag,status FROM sys_holiday_calendar "
                "WHERE tenant_id=100000 AND calendar_code=%s AND deleted=0", (code,))
        self.check(len(calendars) == 1 and calendars[0]["default_flag"] == 0
                   and calendars[0]["status"] == "ENABLED", "HOLIDAY_isolated_nondefault_calendar")
        calendar_id = calendars[0]["calendar_id"]
        case["calendarId"] = calendar_id
        self.save()
        case, task = self.prepare("auth", "HOLIDAY", "HOLIDAY",
            ["日历", "日期", "日期类型", "节假日名称", "是否三倍工资日", "年份", "备注"],
            [[code, "2026-10-01", "法定节假日", "国庆节", "是", 2026, self.state["tag"]]])
        self.check(task["status"] in ("VALIDATED", "SUCCESS") and task["validRows"] == 1
                   and task["invalidRows"] == 0, "HOLIDAY_validated")
        self.execute("auth", case)
        self.result("auth", "HOLIDAY", case)
        saved = self.sql("SELECT date_id,calendar_id,holiday_date,date_type,holiday_name,triple_pay_flag,year_no "
            "FROM sys_holiday_date WHERE tenant_id=100000 AND calendar_id=%s AND deleted=0", (calendar_id,))
        self.check(len(saved) == 1 and saved[0]["holiday_date"].isoformat() == "2026-10-01"
                   and saved[0]["date_type"] == "HOLIDAY" and saved[0]["holiday_name"] == "国庆节"
                   and saved[0]["triple_pay_flag"] == 1 and saved[0]["year_no"] == 2026,
                   "HOLIDAY_exactly_one_row_and_resolved_calendar")
        dates = api(self.client, "GET", "/auth/system/holiday/date/list",
                    params={"calendarId": calendar_id, "yearNo": 2026})
        self.check(len(dates) == 1 and dates[0]["holidayDate"] == "2026-10-01"
                   and dates[0]["triplePayFlag"] == 1, "HOLIDAY_management_api_reads_imported_date")
        self.check(self.default_calendar_snapshot() == case["defaultBaseline"], "HOLIDAY_official_calendar_unchanged")
        case.update(businessRows=saved, complete=True)
        self.save()

    def normal(self, biz):
        if biz == "OPPORTUNITY":
            self.invalid_opportunity_owner()
        tag = self.state["tag"] + "_" + biz
        if biz == "ASSET":
            service, headers, rows = "oa", ["资产编码", "资产名称", "资产分类", "计提折旧", "资产原值"], [
                [tag, "导入验收资产" + tag, "F-ASSET-01", "否", "0.00"]]
            sql = "SELECT asset_id,category_id,depreciation_flag FROM oa_asset WHERE tenant_id=100000 AND asset_code=%s AND deleted=0"
        elif biz == "EMPLOYEE":
            today = self.sql("SELECT CURDATE() AS today")[0]["today"].isoformat()
            service, headers, rows = "hr", ["工号", "姓名", "性别", "入职日期"], [[tag, "导入验收员工" + tag, "男", today]]
            sql = "SELECT id,user_id,employee_status FROM hr_employee WHERE tenant_id=100000 AND employee_no=%s AND deleted=0"
        else:
            case = self.state["cases"].setdefault(biz, {})
            if "owner" not in case:
                owners = self.sql("SELECT e.user_id,e.employee_no FROM hr_employee e "
                    "JOIN sys_user u ON u.user_id=e.user_id AND u.tenant_id=e.tenant_id "
                    "JOIN sys_dept d ON d.dept_id=e.dept_id AND d.tenant_id=e.tenant_id "
                    "WHERE e.tenant_id=100000 AND e.deleted=0 AND e.employee_status IN ('REGULAR','PROBATION') "
                    "AND u.deleted=0 AND u.status='0' AND d.deleted=0 AND d.status='0' "
                    "AND (SELECT COUNT(*) FROM hr_employee other WHERE other.tenant_id=e.tenant_id "
                    "AND other.user_id=e.user_id AND other.deleted=0)=1 ORDER BY e.user_id LIMIT 1")
                if not owners:
                    raise AssertionError("No enabled unique active employee available for opportunity owner")
                case["owner"] = owners[0]
                self.save()
            if case.get("lastTask", {}).get("status") == "FAILED" and case["lastTask"].get("failRows"):
                self.check(not self.sql("SELECT opportunity_id FROM crm_opportunity WHERE tenant_id=100000 "
                    "AND opportunity_name=%s AND deleted=0", (tag,)), "OPPORTUNITY_failed_attempt_zero_business_rows")
                previous = dict(case)
                case = {"owner": previous["owner"], "previousAttempts": [previous]}
                self.state["cases"][biz] = case
                self.save()
            service, headers, rows = "crm", ["操作", "客户", "商机名称", "负责人", "预计金额", "备注"], [
                ["新增", "CUS-T00-001", tag, case["owner"]["employee_no"], "0.00", "普通导入验收" + tag]]
            sql = "SELECT opportunity_id,customer_id,owner_id FROM crm_opportunity WHERE tenant_id=100000 AND opportunity_name=%s AND deleted=0"
        case, task = self.prepare(service, biz, biz, headers, rows)
        self.check(task["status"] in ("VALIDATED", "SUCCESS") and task["validRows"] == 1
                   and task["invalidRows"] == 0, biz + "_validated")
        self.execute(service, case)
        self.result(service, biz, case)
        saved = self.sql(sql, (tag,))
        self.check(len(saved) == 1, biz + "_exactly_one_business_row")
        if biz == "ASSET":
            self.check(saved[0]["category_id"] == 165001 and saved[0]["depreciation_flag"] == 0,
                       "ASSET_reference_resolved_and_no_depreciation")
        elif biz == "EMPLOYEE":
            self.check(saved[0]["user_id"] is None, "EMPLOYEE_did_not_create_login_account")
        else:
            self.check(saved[0]["customer_id"] == 2800001 and saved[0]["owner_id"] == case["owner"]["user_id"],
                       "OPPORTUNITY_customer_and_owner_resolved")
        case.update(businessRows=saved, complete=True)
        self.save()

    def invalid_opportunity_owner(self):
        key = "OPPORTUNITY_INVALID_OWNER"
        if self.state["cases"].get(key, {}).get("complete"):
            return
        bindings = self.sql("SELECT id FROM hr_employee WHERE tenant_id=100000 AND user_id=%s AND deleted=0",
                            (self.state["operatorId"],))
        if len(bindings) <= 1:
            raise AssertionError("Invalid-owner check requires the existing ambiguous operator binding")
        name = self.state["tag"] + "_INVALID_OWNER"
        case, task = self.prepare("crm", "OPPORTUNITY", key,
            ["操作", "客户", "商机名称", "预计金额"], [["新增", "CUS-T00-001", name, "0.00"]])
        self.check(task["status"] == "VALIDATED" and task["validRows"] == 0
                   and task["invalidRows"] == 1, "OPPORTUNITY_invalid_owner_rejected_during_validation")
        preview = api(self.client, "GET", f"/crm/system/import/{case['taskId']}/preview")
        self.check(any(row["rowNo"] == 2 and "唯一在职员工" in row["errorMessage"]
                       for row in preview["errors"]["rows"]), "OPPORTUNITY_invalid_owner_explained")
        self.check(not self.sql("SELECT opportunity_id FROM crm_opportunity WHERE tenant_id=100000 "
            "AND opportunity_name=%s AND deleted=0", (name,)), "OPPORTUNITY_invalid_owner_zero_business_rows")
        case["complete"] = True
        self.save()

    def suppliers(self):
        tag = self.state["tag"]
        headers = ["供应商编号", "供应商名称", "状态"]
        previous = self.state["cases"].get("SUPPLIER", {})
        if previous.get("corrected"):
            self.supplier_duplicates(previous, headers)
            return
        rows = [[tag + "_S1", "导入验收供应商1" + tag, "停用"], [tag + "_S2", "", "停用"]]
        case, task = self.prepare("oa", "SUPPLIER", "SUPPLIER", headers, rows)
        self.check(task["validRows"] == 1 and task["invalidRows"] == 1, "SUPPLIER_one_invalid_row")
        preview = api(self.client, "GET", f"/oa/system/import/{case['taskId']}/preview")
        errors = preview["errors"]["rows"]
        self.check(any(row["rowNo"] == 3 and row["columnName"] == "供应商名称" and row["errorMessage"] for row in errors)
                   and bool(preview["errorCounts"]), "SUPPLIER_error_row_column_and_summary")
        if not case.get("errorFileChecked"):
            deadline = time.monotonic() + 15
            while True:
                report = request(self.client, "GET", f"/oa/system/import/{case['taskId']}/errors/file")
                if report.status_code == 200 and report.content.startswith(b"PK"):
                    break
                if time.monotonic() >= deadline:
                    raise AssertionError("SUPPLIER error workbook not available within 15 seconds")
                time.sleep(.5)
            book = openpyxl.load_workbook(io.BytesIO(report.content))
            sheet = book.worksheets[0]
            self.check(sheet.cell(1, sheet.max_column).value == "错误说明"
                       and sheet.cell(3, sheet.max_column).value
                       and sheet.cell(3, 2).font.color is not None
                       and sheet.cell(3, 2).font.color.type == "indexed"
                       and sheet.cell(3, 2).font.color.indexed == 16, "SUPPLIER_error_workbook_marked_and_explained")
            book.close()
            case["errorFileChecked"] = True
            self.save()
        select = "SELECT supplier_id,supplier_code,status FROM oa_supplier WHERE tenant_id=100000 AND supplier_code IN (%s,%s) AND deleted=0"
        if not case.get("zeroWriteChecked"):
            self.check(not self.sql(select, (tag + "_S1", tag + "_S2")), "SUPPLIER_before_execute_zero_rows")
            denied = request(self.client, "POST", f"/oa/system/import/{case['taskId']}/execute", json={"skipInvalidRows": False})
            body = denied.json()
            self.check((denied.status_code >= 400 or body.get("code") != 200)
                       and "未写入任何数据" in str(body.get("msg", "")), "SUPPLIER_invalid_execute_refused")
            self.check(not self.sql(select, (tag + "_S1", tag + "_S2")), "SUPPLIER_refusal_zero_business_writes")
            case["zeroWriteChecked"] = True
            self.save()
        if not case.get("skipSubmitted") and self.read("oa", case["taskId"])["status"] == "VALIDATED":
            # The intentional invalid execute and skip execute share a 5-second repeat guard.
            time.sleep(5.2)
        self.execute("oa", case, skip=True)
        case["skipSubmitted"] = True
        self.save()
        self.result("oa", "SUPPLIER", case, expected="PARTIAL", total=2, success=1, fail=1)
        saved = self.sql(select, (tag + "_S1", tag + "_S2"))
        if not case.get("corrected"):
            self.check(len(saved) == 1 and saved[0]["supplier_code"] == tag + "_S1"
                       and saved[0]["status"] == "DISABLED", "SUPPLIER_skipped_invalid_only_one_disabled_row")
            time.sleep(5.2)
            denied = request(self.client, "POST", f"/oa/system/import/{case['taskId']}/execute", json={"skipInvalidRows": True})
            body = denied.json()
            self.check((denied.status_code >= 400 or body.get("code") != 200)
                       and "重复" in str(body.get("msg", "")), "SUPPLIER_completed_task_cannot_execute_again")
            self.check(len(self.sql(select, (tag + "_S1", tag + "_S2"))) == 1, "SUPPLIER_repeat_did_not_add_row")
            corrected, checked = self.prepare("oa", "SUPPLIER", "SUPPLIER_CORRECTED", headers,
                [[tag + "_S2", "导入验收供应商2" + tag, "停用"]])
            self.check(checked["validRows"] == 1 and checked["invalidRows"] == 0, "SUPPLIER_corrected_reupload_validated")
            self.execute("oa", corrected)
            self.result("oa", "SUPPLIER", corrected)
            self.check(len(self.sql(select, (tag + "_S1", tag + "_S2"))) == 2, "SUPPLIER_corrected_only_missing_row_added")
            case["corrected"] = True
            self.save()
        self.supplier_duplicates(case, headers)

    def supplier_duplicates(self, case, headers):
        tag = self.state["tag"]
        duplicate, checked = self.prepare("oa", "SUPPLIER", "SUPPLIER_DUPLICATES", headers, [
            [tag + "_S1", "库内重复" + tag, "停用"],
            [tag + "_S3", "文件重复1" + tag, "停用"],
            [tag + "_S3", "文件重复2" + tag, "停用"]])
        duplicate_preview = api(self.client, "GET", f"/oa/system/import/{duplicate['taskId']}/preview")
        duplicate_errors = duplicate_preview["errors"]["rows"]
        self.check(checked["invalidRows"] >= 2 and any(row["rowNo"] == 2 for row in duplicate_errors)
                   and any(row["rowNo"] in (3, 4) for row in duplicate_errors), "SUPPLIER_database_and_file_duplicates_detected")
        case.update(businessRows=self.sql("SELECT supplier_id,supplier_code,status FROM oa_supplier "
            "WHERE tenant_id=100000 AND supplier_code IN (%s,%s) AND deleted=0", (tag + "_S1", tag + "_S2")), complete=True)
        self.save()

    def run(self, businesses):
        ordinary.authenticate(self.client, os.environ["CLOUDFLOW_TEST_ADMIN_USERNAME"], os.environ["CLOUDFLOW_TEST_ADMIN_PASSWORD"])
        info = api(self.client, "GET", "/auth/info")
        self.state["operatorId"] = info["user"]["userId"]
        self.save()
        failures = []
        for biz in businesses:
            if self.state["cases"].get(biz, {}).get("complete"):
                print("SKIP completed", biz, flush=True)
                continue
            try:
                if biz == "SUPPLIER":
                    self.suppliers()
                elif biz == "HOLIDAY":
                    self.holidays()
                else:
                    self.normal(biz)
                print("BUSINESS_COMPLETE", biz, flush=True)
            except Exception as error:
                failures.append(biz)
                self.state["cases"].setdefault(biz, {})["lastError"] = str(error)
                self.save()
                print("FAIL", biz, str(error), flush=True)
        print("RESULT", json.dumps({"state": str(self.path), "tag": self.state["tag"],
            "passed": len(self.state["passed"]), "failed": failures}, ensure_ascii=False), flush=True)
        return 1 if failures else 0

    def close(self):
        try:
            if "Authorization" in self.client.headers:
                api(self.client, "POST", "/auth/logout")
        finally:
            self.client.close()
            self.connection.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--business", nargs="+", choices=BUSINESSES, default=list(BUSINESSES))
    parser.add_argument("--resume", type=Path)
    args = parser.parse_args()
    state_path = args.resume or ROOT / ".cloudflow-runtime" / "server-deployment-20261018" / (
        "business-imports-" + time.strftime("%Y%m%d-%H%M%S") + ".json")
    checks = Checks(state_path)
    try:
        return checks.run(args.business)
    finally:
        checks.close()


if __name__ == "__main__":
    raise SystemExit(main())
